# -*- coding: utf-8 -*-
"""
Script tạo Giáo trình Thực chiến: 
HƯỚNG DẪN TOÀN DIỆN VIBECODING CÙNG ANTIGRAVITY - TỪ Ý TƯỞNG ĐẾN PRODUCTION
Case Study Thực Tế: IELTS Web (Nền Tảng Luyện Thi Chuẩn CDI)
Dành cho Sinh viên & Người mới bắt đầu (Zero-IT to Full-Stack Deploy)
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_run_font(run, font_name="Calibri"):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_border(cell, left_color="2563EB", left_sz="36", bg_color="F8FAFC"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{left_sz}" w:space="0" w:color="{left_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)

def add_header_footer(doc):
    for s in doc.sections:
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("GIÁO TRÌNH THỰC CHIẾN VIBECODING | ANTIGRAVITY AGENT WORKFLOW")
        set_run_font(hrun, "Calibri")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Tài liệu hướng dẫn sinh viên • Dự án IELTS Web (Case Study Đầy Đủ)")
        set_run_font(frun, "Calibri")
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def add_callout(doc, text, title="MẸO HAY CHO SINH VIÊN", box_type="tip"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)

    if box_type == "tip":
        border_col, bg_col = "10B981", "F0FDF4"
        icon = "💡"
    elif box_type == "warning":
        border_col, bg_col = "EF4444", "FEF2F2"
        icon = "⚠️"
    elif box_type == "concept":
        border_col, bg_col = "8B5CF6", "F5F3FF"
        icon = "🧠"
    elif box_type == "tech":
        border_col, bg_col = "0D9488", "F0FDFA"
        icon = "⚙️"
    else:
        border_col, bg_col = "3B82F6", "EFF6FF"
        icon = "📌"

    set_cell_border(cell, left_color=border_col, left_sz="36", bg_color=bg_col)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"{icon} {title}\n")
    set_run_font(r_title, "Calibri")
    r_title.bold = True
    r_title.font.size = Pt(10.5)
    if box_type == "tip":
        r_title.font.color.rgb = RGBColor(4, 120, 87)
    elif box_type == "warning":
        r_title.font.color.rgb = RGBColor(185, 28, 28)
    elif box_type == "concept":
        r_title.font.color.rgb = RGBColor(109, 40, 217)
    elif box_type == "tech":
        r_title.font.color.rgb = RGBColor(15, 118, 110)
    else:
        r_title.font.color.rgb = RGBColor(29, 78, 216)

    r_text = p.add_run(text)
    set_run_font(r_text, "Calibri")
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_prompt_box(doc, prompt_content, prompt_title="PROMPT THỰC HÀNH 1-CHẠM", notes=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)

    set_cell_border(cell, left_color="2563EB", left_sz="36", bg_color="F8FAFC")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r_title = p.add_run(f"🚀 {prompt_title}\n")
    set_run_font(r_title, "Calibri")
    r_title.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    r_guide = p.add_run("(Copy toàn bộ đoạn văn bản bên dưới dán vào ô chat Antigravity):\n\n")
    set_run_font(r_guide, "Calibri")
    r_guide.italic = True
    r_guide.font.size = Pt(9.5)
    r_guide.font.color.rgb = RGBColor(100, 116, 139)

    r_code = p.add_run(prompt_content.strip())
    set_run_font(r_code, "Consolas")
    r_code.font.size = Pt(9.5)
    r_code.font.color.rgb = RGBColor(15, 23, 42)

    if notes:
        r_note = p.add_run(f"\n\n👉 Lưu ý kỹ thuật: {notes}")
        set_run_font(r_note, "Calibri")
        r_note.font.size = Pt(9.5)
        r_note.font.color.rgb = RGBColor(180, 83, 9)
        r_note.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_shading(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        set_run_font(run, "Calibri")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.add_row()
        row_cells = row.cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row_data):
            row_cells[i].width = Inches(col_widths[i])
            set_cell_shading(row_cells[i], bg_col)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=140, right=140)
            
            tcPr = row_cells[i]._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            set_run_font(run, "Calibri")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(51, 65, 85)

def add_h1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_run_font(run, "Calibri")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(30, 58, 138)
    return h

def add_h2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_run_font(run, "Calibri")
    run.bold = True
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(15, 118, 110)
    return h

def add_h3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_run_font(run, "Calibri")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(71, 85, 105)
    return h

def add_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, "Calibri")
        r_pre.bold = True
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    set_run_font(r, "Calibri")
    r.font.size = Pt(10.5)
    r.italic = italic
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, "Calibri")
        r_pre.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    set_run_font(r, "Calibri")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def build_vibecoding_master_doc():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ==========================================
    # TRANG BÌA (COVER PAGE)
    # ==========================================
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.columns[0].width = Inches(6.5)
    c_cell = cover_table.cell(0, 0)
    set_cell_shading(c_cell, "0F172A") # Slate 900 Premium Dark
    set_cell_margins(c_cell, top=550, bottom=550, left=350, right=350)

    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(30)
    cp.paragraph_format.space_after = Pt(10)

    run_tag = cp.add_run("GIÁO TRÌNH HƯỚNG DẪN SINH VIÊN KỶ NGUYÊN AI AGENT\n\n")
    set_run_font(run_tag, "Calibri")
    run_tag.bold = True
    run_tag.font.size = Pt(11)
    run_tag.font.color.rgb = RGBColor(56, 189, 248) # Sky 400

    run_title = cp.add_run("VIBECODING TOÀN DIỆN VỚI ANTIGRAVITY\n")
    set_run_font(run_title, "Calibri")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(255, 255, 255)

    run_sub = cp.add_run("Từ Ý Tưởng, Kiến Trúc, Triển Khai Đến Cloud Database, CI/CD & Production\n\n")
    set_run_font(run_sub, "Calibri")
    run_sub.font.size = Pt(12.5)
    run_sub.font.color.rgb = RGBColor(226, 232, 240)

    cp2 = c_cell.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_before = Pt(15)
    cp2.paragraph_format.space_after = Pt(30)

    run_case = cp2.add_run("CASE STUDY THỰC NGHIỆM: DỰ ÁN IELTS WEB (CHUẨN PHÒNG THI MÁY CDI)\n")
    set_run_font(run_case, "Calibri")
    run_case.bold = True
    run_case.font.size = Pt(11.5)
    run_case.font.color.rgb = RGBColor(250, 204, 21) # Amber Yellow

    run_meta = cp2.add_run(
        "Công nghệ: Google Antigravity + React/Vite + Supabase + Vercel + GitBook + GitHub Actions\n"
        "Đối tượng: Sinh viên mọi chuyên ngành (Không yêu cầu nền tảng IT từ trước)\n"
        "Phiên bản: 2.0 (2026) • Lưu hành nội bộ nghiên cứu & học tập"
    )
    set_run_font(run_meta, "Calibri")
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_page_break()
    add_header_footer(doc)

    # ==========================================
    # LỜI TỰA & TRIẾT LÝ VIBECODING
    # ==========================================
    add_h1(doc, "LỜI TỰA: TRIẾT LÝ VIBECODING TRONG KỶ NGUYÊN AI AGENT")
    add_p(doc, 
        "\"Vibecoding\" là thuật ngữ được giới công nghệ thế giới định nghĩa cho phương thức phát triển phần mềm mới: "
        "con người không còn phải ghi nhớ từng dòng lệnh cú pháp phức tạp, mà tập trung vào cảm thụ sản phẩm (vibe), "
        "đặc tả logic nghiệp vụ và chỉ đạo các AI Agent thông minh (như Google Antigravity) tự tay lập trình.")

    add_p(doc,
        "Tuy nhiên, đối với sinh viên, Vibecoding KHÔNG CÓ NGHĨA là \"nhắm mắt bấm nút bừa bãi\". Một sinh viên làm chủ "
        "Vibecoding thực thụ phải nắm được toàn bộ dòng chảy (pipeline) của một sản phẩm phần mềm hiện đại:")

    add_bullet(doc, "Hiểu được bài toán thực tế, chia nhỏ tính năng thành các bản đặc tả PRD mà AI hiểu được 100%.", "1. Tư duy kiến trúc (Architect): ")
    add_bullet(doc, "Biết cách cấu hình máy tính cá nhân để Antigravity có đầy đủ quyền hạn thực thi lệnh, đọc ghi file và sửa lỗi.", "2. Môi trường phát triển (Dev Environment): ")
    add_bullet(doc, "Hiểu được cách dữ liệu di chuyển từ tệp tĩnh cục bộ (LocalStorage) lên đám mây (Supabase) và phân quyền bảo mật.", "3. Quản trị dữ liệu (Data & Security): ")
    add_bullet(doc, "Biết cách đưa mã nguồn lên GitHub, kết nối kiểm thử tự động CI/CD, đưa web lên Vercel và xuất bản tài liệu GitBook.", "4. Quy trình đóng gói công nghiệp (DevOps & CI/CD): ")

    add_callout(doc,
        "Giáo trình này được đúc kết 100% từ quá trình thực chiến xây dựng dự án IELTS Web đang chạy thực tế trên máy tính của bạn. "
        "Mỗi bước đều có mã nguồn thật, lệnh thật, prompt thật và giải thích cặn kẽ để bất kỳ sinh viên nào cũng có thể "
        "làm theo thành công ngay trong lần đầu tiên.",
        title="MỤC TIÊU CỦA TÀI LIỆU NÀY",
        box_type="concept"
    )

    doc.add_page_break()

    # ==========================================
    # PHẦN 1: CÀI ĐẶT MÔI TRƯỜNG MÁY TÍNH
    # ==========================================
    add_h1(doc, "PHẦN I: THIẾT LẬP MÁY TÍNH ĐỂ VIBECODE VỚI ANTIGRAVITY DỄ DÀNG NHẤT")

    add_p(doc, 
        "Để Antigravity có thể tự động viết code, cài đặt thư viện và khởi chạy máy chủ thay bạn mà không bị chặn quyền, "
        "bạn cần chuẩn bị môi trường máy tính theo 4 bước chuẩn dưới đây.")

    add_h2(doc, "1.1. Cài Đặt Node.js (Runtime Engine)")
    add_p(doc, "Node.js là môi trường bắt buộc để chạy các ứng dụng JavaScript/React hiện đại.")
    add_bullet(doc, "Tải bản LTS tại trang chủ https://nodejs.org.", "Bước 1: ")
    add_bullet(doc, "Cài đặt: Giữ nguyên mọi tùy chọn mặc định, nhấn Next đến khi hoàn tất.", "Bước 2: ")
    add_bullet(doc, "Kiểm tra trong PowerShell bằng lệnh: `node -v` và `npm -v` (nếu hiện phiên bản v20+ là thành công).", "Bước 3: ")

    add_h2(doc, "1.2. Cài Đặt Git Cho Windows (Hệ Thống Quản Lý Mã Nguồn)")
    add_p(doc, "Git là công cụ giúp Antigravity lưu lại các mốc lịch sử dự án, đồng bộ lên GitHub và hoàn tác khi gặp lỗi.")
    add_bullet(doc, "Tải Git từ https://git-scm.com/download/win.", "Bước 1: ")
    add_bullet(doc, "Trong quá trình cài đặt, chọn 'Git from the command line and also from 3rd-party software'.", "Bước 2: ")
    add_bullet(doc, "Cấu hình danh tính lần đầu trong PowerShell:\n`git config --global user.name \"Ten Cua Ban\"`\n`git config --global user.email \"email@gmail.com\"`", "Bước 3: ")

    add_h2(doc, "1.3. Cài Đặt & Cấu Hình Google Antigravity IDE")
    add_p(doc, "Antigravity là IDE thế hệ mới tích hợp trực tiếp AI Agent của Google DeepMind:")
    add_bullet(doc, "Khởi động Antigravity và đăng nhập tài khoản Google.", "Bước 1: ")
    add_bullet(doc, "Mở thư mục làm việc: File -> Open Folder -> Chọn thư mục dự án (ví dụ: `C:\\Projects\\IELTS Web`).", "Bước 2: ")
    add_bullet(doc, "Cấp quyền cho Agent chạy lệnh Terminal: Trong Antigravity, khi Agent đề xuất chạy lệnh (run_command), chọn 'Always allow in this workspace' để AI có thể tự động cài đặt thư viện và chạy máy chủ mà không làm phiền bạn liên tục.", "Bước 3: ")

    add_h2(doc, "1.4. Ma Trận Chọn Model AI (Model Tiering Matrix) Tiết Kiệm Chi Phí")
    add_p(doc, 
        "Một bí quyết sống còn được áp dụng trong file `CODEBASE_MAP.md` của dự án IELTS Web: "
        "Không dùng một model đắt đỏ cho mọi tác vụ! Hãy chuyển đổi model linh hoạt trong Antigravity theo bảng sau:")

    tier_table = doc.add_table(rows=1, cols=4)
    format_table(
        tier_table,
        [1.3, 1.8, 1.6, 1.8],
        ["Tầng công việc", "Loại tác vụ cụ thể", "Model khuyến nghị", "Lợi ích về Token & Tốc độ"],
        [
            ["Tầng 1: Data", "Thêm đề thi, sửa từ vựng, cấu hình tĩnh", "Flash Lite / Mini", "Siêu rẻ, phản hồi tức thì (~1k tokens)"],
            ["Tầng 2: UI Presentation", "Sửa giao diện Tailwind, Icon Lucide, chia cột", "Gemini Flash / Sonnet", "Tốc độ cao, cú pháp React chuẩn (~3-5k tokens)"],
            ["Tầng 3: State & Hooks", "Xử lý phím tắt CDI, đếm giờ, âm thanh", "Gemini Flash", "Tối ưu logic state (~5-8k tokens)"],
            ["Tầng 4: Domain & Algo", "Chấm điểm IELTS, làm tròn TRF, Supabase sync", "Gemini Pro / Opus", "Suy luận sâu, thuật toán chặt chẽ"]
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # ==========================================
    # PHẦN 2: LẬP KẾ HOẠCH & BẢN ĐỒ CODEBASE
    # ==========================================
    add_h1(doc, "PHẦN II: KHỞI TẠO Ý TƯỞNG & THIẾT LẬP BẢN ĐỒ KIẾN TRÚC (CODEBASE MAP)")

    add_p(doc, 
        "Lỗi lớn nhất của sinh viên khi dùng AI là: \"Gõ một câu ngắn ngủn bảo AI làm cả website\". "
        "Cách làm chuyên nghiệp là tạo ra một tài liệu định vị `CODEBASE_MAP.md` để AI Agent chỉ đọc đúng file cần sửa, "
        "tiết kiệm 90% chi phí token và không bao giờ bị lạc lối giữa hàng trăm file.")

    add_h2(doc, "2.1. Phân Tầng Kiến Trúc Dự Án Thực Tế (IELTS Web)")
    add_bullet(doc, "Chứa toàn bộ đề thi `sampleTasks.js`, `readingTasks.js`, `listeningTasks.js`. Khi sửa nội dung bài học, chỉ chạm vào đây.", "Tầng 1 - Dữ liệu tĩnh (`src/data/`): ")
    add_bullet(doc, "Chứa các màn hình luyện thi `WritingWorkspace.jsx`, `SpeakingWorkspace.jsx`, `ReadingWorkspace.jsx` và các Modal tiện ích.", "Tầng 2 - Giao diện (`src/components/`): ")
    add_bullet(doc, "Chứa `appStore.js` (quản lý trạng thái theo mô hình Pub/Sub cực nhẹ), `useAudioEngine.js` (xử lý ghi âm), `useKeyboardShortcuts.js`.", "Tầng 3 - Bộ nhớ ứng dụng (`src/core/` & `src/hooks/`): ")
    add_bullet(doc, "Chứa thuật toán tính điểm chuẩn Cambridge, dịch vụ kết nối đám mây `supabaseClient.js`, công cụ AI chấm bài.", "Tầng 4 - Nghiệp vụ chuyên sâu (`src/services/`): ")

    add_h2(doc, "2.2. Prompt Khởi Tạo Kiến Trúc Dự Án (PRD Generator)")
    
    prompt_arch = """Tôi muốn xây dựng ứng dụng web luyện thi IELTS trên máy tính (CDI Mock Test) theo kiến trúc phân tầng sạch (Clean Architecture).
Công nghệ cốt lõi: React 18, Vite, Tailwind CSS, Lucide Icons.
Hãy đóng vai trò Kỹ sư trưởng (Chief Architect):
1. Phân tích chi tiết yêu cầu bài toán: Luyện 4 kỹ năng Nghe-Nói-Đọc-Viết, đồng hồ đếm ngược, chấm điểm, lưu trữ bài làm.
2. Thiết kế cây thư mục hoàn chỉnh (src/data, src/components, src/core, src/services, src/hooks).
3. Tạo file `CODEBASE_MAP.md` mô tả rõ vai trò từng file và quy tắc nạp context cho AI để tiết kiệm token tối đa.
4. Chưa vội viết code chức năng, hãy trình bày kế hoạch để tôi duyệt trước!"""

    add_prompt_box(doc, prompt_arch, prompt_title="PROMPT THIẾT KẾ KIẾN TRÚC & CODEBASE MAP",
                   notes="Luôn yêu cầu AI lập kế hoạch trước khi code để tránh xung đột mã nguồn sau này.")

    doc.add_page_break()

    # ==========================================
    # PHẦN 3: TRIỂN KHAI FRONTEND TỪNG BƯỚC
    # ==========================================
    add_h1(doc, "PHẦN III: TRIỂN KHAI MÃ NGUỒN FRONTEND TỪNG BƯỚC CÙNG ANTIGRAVITY")

    add_p(doc, 
        "Trong phần này, sinh viên sẽ đồng hành cùng Antigravity để dựng nên các tính năng đặc thù của một phòng thi chuẩn quốc tế.")

    add_h2(doc, "3.1. Dựng Khung Giao Diện Chia Đôi Màn Hình (Split-Screen CDI)")
    add_p(doc, 
        "Kỳ thi máy tính của IDP và British Council luôn sử dụng giao diện 2 cột: đề bài bên trái, bài làm bên phải. "
        "Trong Tailwind CSS, việc này được thực hiện qua cú pháp lưới `grid grid-cols-1 md:grid-cols-2`.")

    prompt_split = """Hãy tạo component `WritingWorkspace.jsx` tại `src/components/`:
1. Chia đôi màn hình 50:50 chuẩn CDI: Cột trái hiển thị đề thi (Task 1 hoặc Task 2), cột phải là Textarea soạn thảo.
2. Thanh công cụ phía trên gồm: Đồng hồ đếm ngược 60 phút, nút Pause/Resume, nút Nộp bài, bộ đếm số từ trực tiếp.
3. Khi người dùng gõ phím, tự động cập nhật số từ và kiểm tra xem đã đủ 150 từ (Task 1) hoặc 250 từ (Task 2) chưa.
4. Tích hợp phím tắt Ctrl+S để lưu nháp tự động và Alt+T để mở thư viện đề thi."""

    add_prompt_box(doc, prompt_split, prompt_title="PROMPT DỰNG PHÒNG THI VIẾT LUẬN CHIA ĐÔI")

    add_h2(doc, "3.2. Quản Lý Trạng Thái Nhẹ Nhàng Bằng appStore.js (Không Cần Redux Cồng Kềnh)")
    add_p(doc, 
        "Nhiều sinh viên bị rối khi học Redux. Trong dự án IELTS Web, Antigravity sử dụng một kho lưu trữ cực nhẹ "
        "dựa trên mẫu thiết kế Pub/Sub (Publish - Subscribe) trong `src/core/appStore.js`:")

    store_code = """// src/core/appStore.js - Quản lý trạng thái toàn cục siêu nhẹ
class AppStore {
  constructor() {
    this.state = {
      activeTab: 'writing',
      currentTask: null,
      timerSeconds: 3600,
      userDraft: '',
      isSubmitted: false
    };
    this.listeners = new Set();
  }
  getState() { return this.state; }
  setState(updater) {
    this.state = typeof updater === 'function' ? updater(this.state) : { ...this.state, ...updater };
    this.listeners.forEach(fn => fn(this.state));
  }
  subscribe(fn) {
    this.listeners.add(fn);
    return () => this.listeners.delete(fn);
  }
}
export const appStore = new AppStore();"""

    add_prompt_box(doc, store_code, prompt_title="CẤU TRÚC KHO LƯU TRỮ TRẠNG THÁI (src/core/appStore.js)",
                   notes="Dễ hiểu hơn Redux 10 lần nhưng xử lý mượt mà việc chia sẻ dữ liệu giữa các màn hình.")

    add_h2(doc, "3.3. Tích Hợp Audio Engine & AI Giám Khảo Phỏng Vấn (Speaking)")
    add_p(doc, 
        "Phần thi Speaking được xử lý qua 2 Custom Hook: `useAudioEngine.js` (quản lý AudioContext và sóng âm Visualizer) "
        "và `useSpeechEngine.js` (giao tiếp Web Speech API). Hãy ra lệnh cho Antigravity tích hợp:")

    prompt_speaking = """Hãy tạo component `SpeakingWorkspace.jsx` kết nối với `useSpeechEngine.js`:
1. Giám khảo AI tự động đọc câu hỏi phỏng vấn bằng giọng bản xứ chuẩn (Text-to-Speech).
2. Khi giám khảo nói xong, tự động mở micro và hiển thị sóng âm động (Audio Visualizer canvas).
3. Chuyển đổi giọng nói của thí sinh thành văn bản (Speech-to-Text) hiển thị theo thời gian thực.
4. Đếm thời gian chuẩn bị 1 phút và thời gian nói 2 phút cho Speaking Part 2."""

    add_prompt_box(doc, prompt_speaking, prompt_title="PROMPT TÍCH HỢP GIÁM KHẢO NÓI & THU ÂM AI")

    doc.add_page_break()

    # ==========================================
    # PHẦN 4: KIỂM THỬ LOCAL & CHẠY 1-CHẠM
    # ==========================================
    add_h1(doc, "PHẦN IV: KIỂM THỬ CỤC BỘ & TẠO FILE KHỞI ĐỘNG 1-CHẠM")

    add_p(doc, 
        "Một sản phẩm tốt phải được kiểm tra kỹ càng trên máy trước khi đưa lên mạng. Trong dự án IELTS Web, "
        "chúng ta có hệ thống test tự động bằng Vitest và kịch bản mô phỏng DOM.")

    add_h2(doc, "4.1. Chạy Bộ Kiểm Thử Tự Động (Automated Testing)")
    add_p(doc, "Trong file `package.json`, các lệnh kiểm thử đã được cấu hình sẵn:")
    add_bullet(doc, "Chạy toàn bộ bài test logic tính điểm và liên kết: `npm test`", "Lệnh 1: ")
    add_bullet(doc, "Chạy giả lập thao tác người dùng trên giao diện: `npm run test:simulation`", "Lệnh 2: ")

    add_callout(doc,
        "Khi sinh viên nhờ Antigravity code thêm tính năng mới, hãy luôn kết thúc prompt bằng câu: "
        "\"Sau khi viết xong, hãy tự chạy `npm test` để đảm bảo không làm gãy các tính năng cũ\". "
        "Đây là tác phong của kỹ sư phần mềm chuyên nghiệp!",
        title="BÍ QUYẾT BẢO VỆ CHẤT LƯỢNG CODE",
        box_type="tip"
    )

    add_h2(doc, "4.2. Tạo File Kích Hoạt Nhanh Trên Windows (.bat)")
    add_p(doc, 
        "Để người dùng không cần biết gõ lệnh `npm run dev`, dự án có sẵn tệp `Khoi_Dong_IELTS_Web.bat` ở thư mục gốc:")

    bat_code = """@echo off
title IELTS Writing Local Server
echo ======================================================
echo    DANG KHOI DONG IELTS WRITING ASSISTANT LOCAL...
echo ======================================================
cd /d "C:\\Projects\\IELTS Web"

:: Kiem tra neu port 3000 dang chay thi mo thang trinh duyet
netstat -ano | findstr :3000 >nul
if %errorlevel% equ 0 (
    echo Server da dang chay san tren port 3000.
    start http://localhost:3000
    exit
)

:: Neu chua chay, bat dau server va mo trinh duyet sau 3 giay
echo Dang khoi dong may chu Vite...
start /b cmd /c "npm run dev -- --port 3000"
timeout /t 3 /nobreak >nul
start http://localhost:3000
echo Ung dung da duoc mo tren trinh duyet!
pause"""

    add_prompt_box(doc, bat_code, prompt_title="NỘI DUNG FILE KHỞI ĐỘNG 1-CHẠM (Khoi_Dong_IELTS_Web.bat)",
                   notes="Chỉ cần double-click file này là web tự bật lên trên Chrome/Edge.")

    doc.add_page_break()

    # ==========================================
    # PHẦN 5: ĐƯA MÃ NGUỒN LÊN GITHUB
    # ==========================================
    add_h1(doc, "PHẦN V: QUẢN LÝ PHIÊN BẢN VÀ ĐƯA MÃ NGUỒN LÊN GITHUB")

    add_p(doc, 
        "GitHub là kho lưu trữ mã nguồn lớn nhất hành tinh. Đưa dự án lên GitHub giúp bạn lưu trữ an toàn, "
        "làm đẹp CV sinh viên và kết nối tự động với nền tảng máy chủ Vercel.")

    add_h2(doc, "5.1. Tạo File .gitignore Chuẩn Mực")
    add_p(doc, "Tuyệt đối không đẩy các thư mục nặng như `node_modules` hay file chứa mật khẩu `.env` lên GitHub:")

    gitignore_code = """# Bỏ qua thư mục thư viện nặng
node_modules/
dist/

# Bỏ qua file chứa khóa bí mật
.env
.env.local

# Bỏ qua file log
npm-debug.log*
yarn-debug.log*"""
    add_prompt_box(doc, gitignore_code, prompt_title="CẤU HÌNH FILE .gitignore CHUẨN")

    add_h2(doc, "5.2. Các Bước Đưa Code Lên GitHub Bằng Lệnh")
    add_bullet(doc, "Truy cập https://github.com/new và tạo repository mới (ví dụ: `ielts-web-cdi`).", "Bước 1: ")
    add_bullet(doc, "Yêu cầu Antigravity đẩy code bằng prompt dưới đây.", "Bước 2: ")

    prompt_git = """Hãy giúp tôi khởi tạo git và đẩy toàn bộ mã nguồn lên GitHub:
1. Chạy lệnh `git init` và tạo nhánh chính `main`.
2. Commit toàn bộ file với thông điệp: 'feat: initial release of IELTS CDI Web platform'.
3. Kết nối với remote repository của tôi tại địa chỉ [DÁN_LINK_GITHUB_CỦA_BẠN_VÀO_ĐÂY].
4. Đẩy mã nguồn lên bằng lệnh `git push -u origin main`."""

    add_prompt_box(doc, prompt_git, prompt_title="PROMPT ĐẨY MÃ NGUỒN LÊN GITHUB TỰ ĐỘNG")

    doc.add_page_break()

    # ==========================================
    # PHẦN 6: KẾT NỐI DATABASE SUPABASE
    # ==========================================
    add_h1(doc, "PHẦN VI: KẾT NỐI CƠ SỞ DỮ LIỆU ĐÁM MÂY SUPABASE & BẢO MẬT RLS")

    add_p(doc, 
        "Supabase là nền tảng Backend-as-a-Service mã nguồn mở mạnh mẽ nhất hiện nay, cung cấp cơ sở dữ liệu PostgreSQL "
        "cực nhanh và hệ thống đăng nhập tài khoản an toàn.")

    add_h2(doc, "6.1. Cấu Trúc Bảng Dữ Liệu Thực Tế (supabase_schema.sql)")
    add_p(doc, "Trong dự án IELTS Web, file `supabase_schema.sql` định nghĩa 3 bảng dữ liệu cốt lõi:")
    add_bullet(doc, "Lưu trữ nội dung bài viết luận của thí sinh, số từ, thời gian hoàn thành và kết quả chấm điểm JSON.", "1. user_submissions: ")
    add_bullet(doc, "Sổ tay từ vựng cá nhân, nghĩa tiếng Việt, câu ví dụ và chủ đề.", "2. user_vocab: ")
    add_bullet(doc, "Thư viện đề thi riêng tư do người dùng tự nạp hoặc đề cộng đồng chia sẻ.", "3. user_custom_tasks: ")

    add_h2(doc, "6.2. Bảo Mật Cấp Hàng (Row Level Security - RLS)")
    add_callout(doc,
        "Tại sao RLS quan trọng đối với sinh viên?\n"
        "Nếu không có RLS, bất kỳ ai cũng có thể đọc trộm hoặc xóa bài làm của sinh viên khác! "
        "Trong `supabase_schema.sql`, chúng ta bật chính sách RLS:\n"
        "`CREATE POLICY \"Users can only read their own submissions\" ON public.user_submissions FOR SELECT USING (auth.uid() = user_id);`\n"
        "Chỉ có chính chủ sở hữu tài khoản mới có quyền xem và sửa bài viết của mình.",
        title="BẢO MẬT DỮ LIỆU CHUẨN CÔNG NGHIỆP",
        box_type="warning"
    )

    add_h2(doc, "6.3. Cấu Hình Khóa Kết Nối Biến Môi Trường (.env)")
    add_p(doc, "Tạo file `.env` tại thư mục gốc với 2 khóa lấy từ Supabase Dashboard (Settings -> API):")
    env_sample = """VITE_SUPABASE_URL=https://your-project-id.supabase.co
VITE_SUPABASE_ANON_KEY=eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9..."""
    add_prompt_box(doc, env_sample, prompt_title="CẤU HÌNH BIẾN MÔI TRƯỜNG (.env)")

    doc.add_page_break()

    # ==========================================
    # PHẦN 7: TRIỂN KHAI LÊN VERCEL APP
    # ==========================================
    add_h1(doc, "PHẦN VII: ĐỒNG BỘ VÀ XUẤT BẢN RA INTERNET VỚI VERCEL APP")

    add_p(doc, 
        "Vercel là dịch vụ lưu trữ đám mây số 1 cho các ứng dụng React. Mỗi khi bạn đẩy code lên GitHub, "
        "Vercel sẽ tự động build lại trang web và cập nhật trong 10 giây!")

    add_h2(doc, "7.1. Cấu Hình SPA Rewrite Trong vercel.json (Tránh Lỗi 404)")
    add_p(doc, 
        "Ứng dụng React là Single Page Application (SPA). Nếu không cấu hình định tuyến, khi người dùng bấm F5 tại trang "
        "`/reading` hoặc `/speaking`, web sẽ bị báo lỗi '404 Not Found'. File `vercel.json` giải quyết triệt để vấn đề này:")

    vercel_code = """{
  "rewrites": [
    { "source": "/(.*)", "destination": "/" }
  ]
}"""
    add_prompt_box(doc, vercel_code, prompt_title="CẤU HÌNH vercel.json CHỐNG LỖI 404 KHI RELOAD")

    add_h2(doc, "7.2. Các Bước Đưa Lên Vercel Trong 2 Phút")
    add_bullet(doc, "Truy cập https://vercel.com và bấm 'Add New' -> 'Project'.", "Bước 1: ")
    add_bullet(doc, "Chọn repository GitHub của dự án vừa tạo.", "Bước 2: ")
    add_bullet(doc, "Tại mục 'Environment Variables', dán 2 biến `VITE_SUPABASE_URL` và `VITE_SUPABASE_ANON_KEY`.", "Bước 3: ")
    add_bullet(doc, "Nhấn 'Deploy'. Chờ 1 phút để nhận đường link web trực tuyến dạng: `https://ten-du-an.vercel.app`.", "Bước 4: ")

    doc.add_page_break()

    # ==========================================
    # PHẦN 8: XÂY DỰNG TÀI LIỆU VỚI GITBOOK & CI/CD
    # ==========================================
    add_h1(doc, "PHẦN VIII: TỰ ĐỘNG HÓA TÀI LIỆU VỚI GITBOOK & GITHUB ACTIONS")

    add_p(doc, 
        "Một sản phẩm phần mềm chuyên nghiệp luôn đi kèm tài liệu hướng dẫn tra cứu đầy đủ. "
        "Dự án IELTS Web sở hữu cả một thư viện cẩm nang khổng lồ (IELTS Master Handbook) trong thư mục `docs/` "
        "và được tự động đồng bộ hóa lên GitBook.")

    add_h2(doc, "8.1. Cấu Hình GitBook Docs (`gitbook-docs.yaml` & `.gitbook.yaml`)")
    add_p(doc, "File `gitbook-docs.yaml` chỉ định cho GitBook biết thư mục chứa tài liệu Markdown:")

    gitbook_yaml = """$schema: https://api.gitbook.com/gitbook-docs.yaml
site:
  title: IELTS Master Handbook
  structure:
    - type: space
      key: ielts-handbook
      title: IELTS Master Handbook
      path: docs
      default: true
      content:
        directory: ./docs"""
    add_prompt_box(doc, gitbook_yaml, prompt_title="CẤU HÌNH TÀI LIỆU GITBOOK (gitbook-docs.yaml)")

    add_h2(doc, "8.2. Script Tự Động Đồng Bộ Tính Năng (`sync_features_to_gitbook.js`)")
    add_p(doc, 
        "Khi đội ngũ phát triển thêm một tính năng mới trong `src/core/featureRegistry.js`, "
        "script `scripts/sync_features_to_gitbook.js` sẽ tự động quét và cập nhật tài liệu tương ứng trong `docs/`.")

    add_h2(doc, "8.3. Đường Ống CI/CD Tự Động Kiểm Tra Siêu Liên Kết (.github/workflows/sync-docs.yml)")
    add_p(doc, 
        "Dự án sử dụng GitHub Actions để đảm bảo 'Zero Broken Links' (không bao giờ có liên kết chết trong hơn 640 bài viết tài liệu):")

    add_bullet(doc, "Mỗi khi có code mới đẩy lên nhánh `main`, máy chủ Ubuntu của GitHub tự động khởi chạy.", "Tự động kích hoạt: ")
    add_bullet(doc, "Chạy `node tests/test_step89_markdown_links_integrity.js` kiểm tra toàn bộ liên kết nội bộ.", "Kiểm tra toàn vẹn: ")
    add_bullet(doc, "Nếu tài liệu có thay đổi, bot GitHub Actions tự động commit và đẩy bản cập nhật lên GitBook.", "Tự động xuất bản: ")

    doc.add_page_break()

    # ==========================================
    # PHẦN 9: CẨM NANG BÍ QUYẾT VIBECODING
    # ==========================================
    add_h1(doc, "PHẦN IX: CẨM NANG BÍ QUYẾT VIBECODING DÀNH CHO SINH VIÊN")

    add_p(doc, "Để trở thành một sinh viên làm chủ công nghệ xuất sắc, hãy khắc cốt ghi tâm 5 nguyên tắc sau:")

    add_h2(doc, "Nguyên Tắc 1: Chia Nhỏ Tác Vụ (Task Chunking)")
    add_p(doc, "Không bao giờ yêu cầu AI: \"Hãy làm cả 4 kỹ năng Nghe Nói Đọc Viết cho tôi\". Thay vào đó, hãy chia nhỏ:")
    add_bullet(doc, "Buổi 1: Dựng khung sườn giao diện Writing và bộ đếm từ.", "- ")
    add_bullet(doc, "Buổi 2: Tích hợp đồng hồ đếm ngược và nạp 5 đề thi mẫu.", "- ")
    add_bullet(doc, "Buổi 3: Thêm modal chấm điểm và kết nối LocalStorage.", "- ")

    add_h2(doc, "Nguyên Tắc 2: Đọc Diff (Xem AI Đã Sửa Những Gì)")
    add_p(doc, "Trước khi chấp nhận một thay đổi mã nguồn trong Antigravity, hãy nhìn qua cột xem trước (Diff Preview). "
               "Xem AI đã thêm đoạn code nào và xóa đoạn code nào. Điều này giúp bạn học lập trình nhanh gấp 5 lần đọc sách giáo trình thông thường.")

    add_h2(doc, "Nguyên Tắc 3: Xử Lý Lỗi Màn Hình Đỏ Theo Chu trình 3 Bước")
    add_bullet(doc, "Bước 1: Không hoảng sợ, không tự ý sửa tay khi chưa hiểu.", "Bình tĩnh: ")
    add_bullet(doc, "Bước 2: Copy toàn bộ thông báo lỗi từ terminal hoặc thanh công cụ kiểm tra (F12 Console).", "Thu thập log: ")
    add_bullet(doc, "Bước 3: Dán cho Antigravity kèm câu lệnh: \"Hãy giải thích tại sao lỗi này xảy ra và sửa tận gốc giúp tôi\".", "Ra lệnh sửa: ")

    add_h2(doc, "10 Mẫu Prompt Phổ Biến Nhất Cho Sinh Viên Khi Vibecoding")

    prompt_matrix = [
        ["1. Thêm tính năng mới", "Hãy thêm tính năng [X] vào component [Y]. Đảm bảo không làm thay đổi các chức năng đang chạy tốt."],
        ["2. Sửa lỗi giao diện", "Giao diện trên màn hình điện thoại bị tràn chữ. Hãy chỉnh lại Tailwind CSS để responsive chuẩn mobile."],
        ["3. Tối ưu tốc độ", "Trang web tải hơi chậm khi bấm nạp đề. Hãy tối ưu lại useMemo và useCallback để component không render thừa."],
        ["4. Viết bài test", "Hãy viết một file test trong thư mục tests/ để kiểm tra logic tính điểm band score xem đã chính xác chưa."],
        ["5. Giải thích mã nguồn", "Hãy giải thích đoạn code trong file [X] bằng ngôn ngữ đời thường cho người mới bắt đầu dễ hiểu."],
        ["6. Tạo dữ liệu mẫu", "Hãy tạo 20 câu hỏi trắc nghiệm từ vựng chủ đề Environment dưới dạng mảng JSON trong src/data/."],
        ["7. Tích hợp âm thanh", "Hãy thêm âm thanh thông báo vui nhộn khi người dùng hoàn thành bài thi và đạt số từ mục tiêu."],
        ["8. Đổi màu giao diện", "Hãy chuyển theme màu chủ đạo từ Navy Blue sang Dark Mode hiện đại (nền đen #0F172A, chữ trắng)."],
        ["9. Kiểm tra bảo mật", "Hãy rà soát file supabaseClient.js xem các khóa API và truy vấn đã an toàn theo chuẩn RLS chưa."],
        ["10. Viết hướng dẫn", "Hãy viết một file README.md thật đẹp mắt hướng dẫn người dùng cách chạy dự án này trên máy tính của họ."]
    ]

    pm_table = doc.add_table(rows=1, cols=2)
    format_table(
        pm_table,
        [2.2, 4.3],
        ["Tình huống sử dụng", "Mẫu câu lệnh Prompt copy dùng ngay"],
        prompt_matrix
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2(doc, "Lời Kết")
    add_p(doc, 
        "Công nghệ không phải là bức tường ngăn cách, mà là chiếc đòn bẩy vĩ đại. Với Antigravity và tư duy Vibecoding, "
        "bất kỳ sinh viên nào cũng có thể hiện thực hóa những ý tưởng phần mềm phục vụ cộng đồng, học tập và khởi nghiệp. "
        "Chúc các bạn có những trải nghiệm kiến tạo sản phẩm thật thú vị và tự hào!")

    # Lưu tệp
    output_filename = "Giao_Trinh_Thuc_Chien_Vibecoding_Antigravity_IELTS_Web.docx"
    doc.save(output_filename)
    print(f"Master guide created successfully: {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    build_vibecoding_master_doc()
